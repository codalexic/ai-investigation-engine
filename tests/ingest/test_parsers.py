import io
from datetime import date

import pytest
from docx import Document
from pypdf import PdfReader, PdfWriter
from reportlab.pdfgen.canvas import Canvas

from app.core.errors import IngestionError, UnsupportedFormatError
from app.ingest.metadata import parse_date
from app.ingest.parsers import decode, get_parser, normalize_text


def test_pdf():
    stream = io.BytesIO()
    canvas = Canvas(stream)
    canvas.setTitle("Operations")
    canvas.setAuthor("Team")
    canvas.drawString(40, 700, "First page")
    canvas.showPage()
    canvas.drawString(40, 700, "Second page")
    canvas.save()
    parsed = get_parser("x.PDF").parse(stream.getvalue(), "x.PDF")
    assert parsed.metadata["page_count"] == 2
    assert parsed.metadata["title"] == "Operations"
    assert parsed.metadata["author"] == "Team"
    assert parsed.text == "First page\n\nSecond page"
    encrypted = io.BytesIO()
    writer = PdfWriter()
    writer.append(PdfReader(io.BytesIO(stream.getvalue())))
    writer.encrypt("secret")
    writer.write(encrypted)
    with pytest.raises(IngestionError):
        get_parser("x.pdf").parse(encrypted.getvalue(), "x.pdf")
    with pytest.raises(IngestionError):
        get_parser("x.pdf").parse(b"corrupt", "x.pdf")


def test_docx():
    doc = Document()
    doc.core_properties.title = "Handbook"
    doc.core_properties.author = "Team"
    doc.add_paragraph("Preface")
    doc.add_heading("Root", 1)
    doc.add_paragraph("Instructions")
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    doc.add_heading("Child", 2)
    doc.add_paragraph("Details")
    stream = io.BytesIO()
    doc.save(stream)
    parsed = get_parser("x.docx").parse(stream.getvalue(), "x.docx")
    assert parsed.metadata["title"] == "Handbook"
    assert parsed.sections[1].heading == "Root"
    assert parsed.sections[2].level == 2
    assert "A | B" in parsed.sections[1].text
    assert parsed.text.index("A | B") < parsed.text.index("Child")
    with pytest.raises(IngestionError):
        get_parser("x.docx").parse(b"bad", "x.docx")


def test_markdown_frontmatter_encoding_normalization():
    data = (
        b"---\ntitle: Example\ndate: 2024-03-14\n---\n# Root\nCaf\xe9  \n\n\n\n## Child\nBody\x00"
    )
    parsed = get_parser("x.md").parse(data, "x.md")
    assert parsed.metadata["title"] == "Example"
    assert parsed.sections[1].heading == "Child"
    assert "Caf\u00e9" in parsed.text
    assert "\x00" not in parsed.text and "\n\n\n" not in parsed.text
    assert normalize_text("\uff21 \x00  \n\n\nB  ") == "A\n\nB"
    assert decode(b"a\x81") == "a\x81"
    assert get_parser("x.txt").parse(b"hello", "x.txt").text == "hello"


def test_html():
    parsed = get_parser("x.html").parse(
        b"<title>Title</title><nav>NO</nav><script>NO</script><style>NO</style>"
        b"<h1>Root</h1><p>Some <b>bold</b> text</p><h2>Child</h2><p>Body</p>",
        "x.html",
    )
    assert parsed.metadata["title"] == "Title"
    assert "NO" not in parsed.text
    assert [s.heading for s in parsed.sections] == ["Root", "Child"]
    assert "bold" in parsed.text


@pytest.mark.parametrize(
    "filename,data",
    [
        ("x.json", b'[{"text":"Hello", "author":"Team"},{"body":"World"}]'),
        ("x.jsonl", b'{"text":"Hello", "author":"Team"}\n{"body":"World"}'),
    ],
)
def test_json(filename, data):
    parsed = get_parser(filename).parse(data, filename)
    assert len(parsed.sections) == 2
    assert "text: Hello" in parsed.text and "body: World" in parsed.text
    assert parsed.metadata["author"] == "Team"
    assert "text" not in parsed.metadata


def test_json_errors_and_object():
    parser = get_parser("x.json")
    assert parser.parse(b'{"content":"Body"}', "x.json").text == "content: Body"
    for data in (b"broken", b"[3]"):
        with pytest.raises(IngestionError):
            parser.parse(data, "x.json")


def test_csv_and_unsupported():
    parsed = get_parser("x.csv").parse(b'name,date\n"a,b",2024-03-14\nc,2024-03-15', "x.csv")
    assert len(parsed.sections) == 2
    assert parsed.sections[0].text == "name: a,b; date: 2024-03-14"
    assert parse_date(parsed.metadata["created_at"]) == date(2024, 3, 14)
    with pytest.raises(UnsupportedFormatError):
        get_parser("unknown.bin")
