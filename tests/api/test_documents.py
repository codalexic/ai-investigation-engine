import io
import json

import pytest
from docx import Document
from pydantic import SecretStr
from reportlab.pdfgen.canvas import Canvas


def upload(client, text=b"payment-service rollback succeeded", name="incident.txt", **data):
    return client.post("/api/documents", files={"file": (name, text)}, data=data)


@pytest.mark.parametrize("extension", ["txt", "md", "html", "pdf", "docx", "json", "jsonl", "csv"])
def test_formats(api_client, extension):
    data = b"payment-service rollback succeeded"
    if extension == "pdf":
        stream = io.BytesIO()
        canvas = Canvas(stream)
        canvas.drawString(40, 700, data.decode())
        canvas.save()
        data = stream.getvalue()
    elif extension == "docx":
        stream = io.BytesIO()
        document = Document()
        document.add_paragraph(data.decode())
        document.save(stream)
        data = stream.getvalue()
    elif extension in {"json", "jsonl"}:
        data = json.dumps({"text": data.decode()}).encode()
    elif extension == "csv":
        data = b"text,date\npayment-service rollback succeeded,2024-03-01"
    elif extension == "html":
        data = b"<h1>Incident</h1><p>payment-service rollback succeeded</p>"
    response = upload(api_client, data, f"incident.{extension}")
    assert response.status_code == 200, response.text
    assert response.json()["status"] == "created"
    assert response.json()["chunk_count"] >= 1


def test_document_lifecycle(api_client):
    first = upload(
        api_client, doc_type="incident", metadata='{"origin":"seed","dataset":"demo"}'
    ).json()
    identifier = first["doc_id"]
    assert upload(api_client).json()["status"] == "duplicate"
    assert api_client.get("/api/documents", params={"q": "absent"}).json()["total"] == 0
    listing = api_client.get(
        "/api/documents", params={"doc_type": "incident", "origin": "user", "limit": 1}
    ).json()
    assert listing["total"] == 1
    assert listing["items"][0]["origin"] == "user"
    assert api_client.get("/api/documents", params={"offset": 1}).json()["items"] == []
    detail = api_client.get(f"/api/documents/{identifier}").json()
    assert len(detail["versions"]) == 1
    chunks = api_client.get(f"/api/documents/{identifier}/chunks").json()
    assert chunks[0]["text"] == "payment-service rollback succeeded"
    response = api_client.put(
        f"/api/documents/{identifier}",
        files={
            "file": (
                "new.txt",
                b"A completely different replacement document about new deployments",
            )
        },
    )
    assert response.json()["version"] == 2
    assert api_client.get(f"/api/documents/{identifier}").json()["extra"]["dataset"] == "demo"
    assert len(api_client.get(f"/api/documents/{identifier}").json()["versions"]) == 2
    assert api_client.delete(f"/api/documents/{identifier}").status_code == 204
    assert api_client.get(f"/api/documents/{identifier}").status_code == 404
    assert api_client.delete(f"/api/documents/{identifier}").status_code == 404


def test_upload_errors(api_client):
    assert upload(api_client, name="bad.exe").status_code == 415
    assert upload(api_client, b"x" * (1024 * 1024 + 1)).status_code == 413
    assert upload(api_client, metadata="bad").status_code == 422
    assert upload(api_client, metadata="[]").status_code == 422


@pytest.mark.parametrize("prefix", ["hotpot:", "pm:", "syn:", "explicit-"])
def test_seed_protection(api_client, prefix):
    container = api_client.app.state.container
    identifier = prefix + "seed"
    container.pipeline.ingest(
        b"seed content about an indexed corpus",
        "seed.txt",
        identifier,
        {"origin": "seed"} if prefix == "explicit-" else {},
    )
    container.settings.admin_token = SecretStr("secret")
    path = f"/api/documents/{identifier}"
    assert api_client.delete(path).status_code == 403
    assert api_client.put(path, files={"file": ("new.txt", b"changed")}).status_code == 403
    assert api_client.delete(path, headers={"X-Admin-Token": "wrong"}).status_code == 403
    container.settings.admin_token = None
    container.settings.app_env = "prod"
    assert api_client.delete(path).status_code == 403
    container.settings.app_env = "local"
    assert (
        api_client.put(path, files={"file": ("new.txt", b"updated seed content")}).status_code
        == 200
    )
    container.settings.admin_token = SecretStr("secret")
    assert api_client.delete(path, headers={"X-Admin-Token": "secret"}).status_code == 204


def test_near_duplicate(api_client):
    text = " ".join(f"word{i}" for i in range(400))
    upload(api_client, text.encode())
    result = upload(api_client, (text + " additional").encode(), "second.txt").json()
    assert result["status"] == "near_duplicate"
    assert result["duplicate_of"]


@pytest.mark.parametrize("extension", ["docx", "pdf"])
def test_created_metadata_persists_end_to_end(api_client, extension):
    from datetime import datetime

    stream = io.BytesIO()
    if extension == "docx":
        document = Document()
        document.core_properties.created = datetime(2024, 3, 12, 10, 30)
        document.add_paragraph("Document with core creation properties")
        document.save(stream)
    else:
        canvas = Canvas(stream, invariant=True)
        canvas.drawString(40, 700, "Document with PDF creation properties")
        canvas.save()
    response = upload(api_client, stream.getvalue(), "dated." + extension)
    assert response.status_code == 200, response.text
    key = response.json()["doc_id"]
    detail = api_client.get(f"/api/documents/{key}").json()
    assert isinstance(detail["extra"]["created_at"], str)
    assert detail["created_at"]
    assert api_client.get(f"/api/documents/{key}/chunks").json()


def test_metadata_only_replacement_updates_payload_and_version(api_client):
    from app.ingest.registry import IndexMeta

    container = api_client.app.state.container
    first = upload(api_client).json()
    key = first["doc_id"]
    with container.registry.session() as session:
        previous = session.get(IndexMeta, "index_version").value
    response = api_client.put(
        f"/api/documents/{key}",
        files={"file": ("incident.txt", b"payment-service rollback succeeded")},
        data={
            "metadata": json.dumps(
                {"title": "Corrected", "created_at": "2025-03-12", "owner": "Operations"}
            ),
            "doc_type": "runbook",
        },
    )
    assert response.json()["status"] == "metadata_updated"
    detail = api_client.get(f"/api/documents/{key}").json()
    assert detail["title"] == "Corrected" and detail["extra"]["owner"] == "Operations"
    assert len(detail["versions"]) == 1
    payload = api_client.get(f"/api/documents/{key}/chunks").json()[0]
    assert payload["title"] == "Corrected" and payload["doc_type"] == "runbook"
    assert payload["created_at"].startswith("2025-03-12")
    with container.registry.session() as session:
        assert session.get(IndexMeta, "index_version").value == previous + 1


def test_replace_lookup_runs_in_threadpool(api_client, monkeypatch):
    import threading

    first = upload(api_client).json()
    registry = api_client.app.state.container.registry
    original = registry.get
    threads = []

    def get(key):
        threads.append(threading.current_thread().name)
        return original(key)

    monkeypatch.setattr(registry, "get", get)
    response = api_client.put(
        f"/api/documents/{first['doc_id']}", files={"file": ("x.txt", b"replacement")}
    )
    assert response.status_code == 200
    assert threads and all("portal" not in name.lower() for name in threads)
